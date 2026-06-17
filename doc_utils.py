"""Shared utilities for generating Word documents with python-docx."""

from docx.enum.text import WD_ALIGN_PARAGRAPH


def add_heading(doc, text, level=1):
    """Add a left-aligned heading to the document."""
    h = doc.add_heading(text, level=level)
    h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return h


def add_centered_title(doc, title, subtitle=None):
    """Add a center-aligned title (level-0 heading) with an optional subtitle."""
    h = doc.add_heading(title, 0)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_section(doc, title, description, level=1):
    """Add a section heading followed by a description paragraph."""
    add_heading(doc, title, level=level)
    doc.add_paragraph(description)


def add_mermaid_section(doc, title, mermaid_code, description=None, level=2):
    """Add a Mermaid diagram section: heading + optional description + code block."""
    if title:
        add_heading(doc, title, level=level)
    if description:
        doc.add_paragraph(description)
    doc.add_paragraph(mermaid_code)


def add_data_table(doc, headers, rows, style='Table Grid'):
    """Add a table with a header row and data rows.

    Parameters
    ----------
    doc : Document
        The python-docx Document instance.
    headers : list[str]
        Column header labels.
    rows : list[tuple[str, ...]]
        Each tuple contains one cell value per column.
    style : str
        Word table style name.
    """
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = style
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row_data in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_data):
            cells[i].text = value
    return table
