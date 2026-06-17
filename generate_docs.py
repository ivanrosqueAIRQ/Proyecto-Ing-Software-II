import os

from docx import Document

from doc_utils import (
    add_centered_title,
    add_data_table,
    add_heading,
    add_mermaid_section,
    add_section,
)

# ---------------------------------------------------------------------------
# Mermaid diagram definitions
# ---------------------------------------------------------------------------

MERMAID_CLASS_DIAGRAM = """
classDiagram
    class Solicitud {
        -Long id
        -String numeroCuenta
        -TipoTramite tipoTramite
        -String motivo
        -EstadoSolicitud estado
        -String resueltoPor
        -LocalDateTime fechaRegistro
    }
    class SolicitudDAO {
        <<interface>>
        +guardar(Solicitud) Solicitud
        +actualizar(Solicitud) Solicitud
        +buscarPorId(Long) Solicitud
        +listarTodas() List~Solicitud~
    }
    class SolicitudDAOImpl {
        -EntityManager em
    }
    class SolicitudServicio {
        -SolicitudDAO solicitudDAO
        -InvocadorComandos invocador
        +registrar(Solicitud) Solicitud
        +aprobarManual(Solicitud) Solicitud
        +rechazarManual(Solicitud) Solicitud
    }
    class TramiteBean {
        -SolicitudServicio servicio
        -Solicitud nuevaSolicitud
        +registrarTramite()
        +aprobar(Solicitud)
        +rechazar(Solicitud)
    }

    SolicitudDAO <|.. SolicitudDAOImpl
    SolicitudServicio --> SolicitudDAO
    SolicitudServicio --> Solicitud
    TramiteBean --> SolicitudServicio
    TramiteBean --> Solicitud
"""

MERMAID_SEQUENCE_DIAGRAM = """
sequenceDiagram
    actor Usuario
    participant TramiteBean
    participant SolicitudServicio
    participant GeneradorFolios
    participant CadenaResponsabilidad
    participant SolicitudDAO
    participant BaseDeDatos

    Usuario->>TramiteBean: registrarTramite()
    TramiteBean->>SolicitudServicio: registrar(Solicitud)
    SolicitudServicio->>GeneradorFolios: siguienteFolio()
    GeneradorFolios-->>SolicitudServicio: folio
    SolicitudServicio->>CadenaResponsabilidad: procesar(Solicitud)
    CadenaResponsabilidad-->>SolicitudServicio: validacion
    SolicitudServicio->>SolicitudDAO: guardar(Solicitud)
    SolicitudDAO->>BaseDeDatos: persist()
    BaseDeDatos-->>SolicitudDAO: success
    SolicitudDAO-->>SolicitudServicio: Solicitud persistida
    SolicitudServicio-->>TramiteBean: success
    TramiteBean-->>Usuario: Muestra Mensaje (FacesMessage)
"""

MERMAID_ER_DIAGRAM = """
erDiagram
    SOLICITUD_TRAMITE {
        BIGINT ID PK
        VARCHAR NUM_CUENTA
        VARCHAR TIPO_TRAMITE
        VARCHAR MOTIVO
        VARCHAR ESTADO
        VARCHAR RESUELTO_POR
        TIMESTAMP FECHA_REGISTRO
    }
"""

# ---------------------------------------------------------------------------
# Database schema definition
# ---------------------------------------------------------------------------

DB_TABLE_HEADERS = ('Columna', 'Tipo de Dato', 'Nulo', 'Descripcion')

DB_TABLE_ROWS = [
    ('ID', 'BIGINT (PK)', 'No', 'Identificador autoincremental'),
    ('NUM_CUENTA', 'VARCHAR(7)', 'No', 'Numero de cuenta del alumno'),
    ('TIPO_TRAMITE', 'VARCHAR(20)', 'No', 'Tipo de tramite solicitado'),
    ('MOTIVO', 'VARCHAR(255)', 'Si', 'Descripcion o justificacion'),
    ('ESTADO', 'VARCHAR(30)', 'No', 'PENDIENTE, APROBADO, etc.'),
    ('RESUELTO_POR', 'VARCHAR(60)', 'Si', 'Usuario que proceso la solicitud'),
    ('FECHA_REGISTRO', 'TIMESTAMP', 'No', 'Fecha de creacion'),
]

# ---------------------------------------------------------------------------
# Document generation
# ---------------------------------------------------------------------------


def generate_doc():
    doc = Document()

    add_centered_title(
        doc,
        'Documentacion del Proyecto SGSA',
        'Sistema de Gestion de Solicitudes Academicas (Java EE)',
    )

    doc.add_page_break()

    # 1. Introduccion
    add_section(
        doc,
        '1. Introduccion',
        'Este documento contiene la documentacion tecnica del proyecto SGSA, '
        'incluyendo diagramas UML y el modelo de base de datos relacional. '
        'El sistema esta construido con arquitectura N-Capas usando JSF, EJB/CDI y JPA.',
    )

    # 2. Diagramas UML
    add_section(doc, '2. Diagramas UML', '')

    add_mermaid_section(
        doc,
        '2.1 Diagrama de Clases',
        MERMAID_CLASS_DIAGRAM,
        description='Representacion de las capas de Modelo, Controlador y Negocio.',
    )

    add_mermaid_section(
        doc,
        '2.2 Diagrama de Secuencia (Registro de Tramite)',
        MERMAID_SEQUENCE_DIAGRAM,
    )

    doc.add_page_break()

    # 3. Modelo de Base de Datos
    add_section(
        doc,
        '3. Modelo de Base de Datos',
        'El sistema utiliza Apache Derby. La tabla principal es '
        'SOLICITUD_TRAMITE que persiste la entidad Solicitud.',
    )

    add_heading(doc, 'Tabla: SOLICITUD_TRAMITE', level=2)
    add_data_table(doc, DB_TABLE_HEADERS, DB_TABLE_ROWS)

    add_mermaid_section(
        doc, '', MERMAID_ER_DIAGRAM,
        description='\nDiagrama ER (Mermaid):',
    )

    output_path = os.path.join(os.path.dirname(__file__), 'Documentacion_SGSA.docx')
    doc.save(output_path)
    print(f'Documento generado en: {output_path}')


if __name__ == '__main__':
    generate_doc()
